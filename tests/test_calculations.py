"""Synthetic analytics regression checks; no app startup or database access.

Run in the app container after copying tests into /app/tests:
    docker exec mutuals_tracker_app python -m unittest discover -s tests -v
Or with dependencies installed: python -m unittest discover -s tests
"""
import ast
import asyncio
import copy
import difflib
import io
import json
import logging
import math
import os
from pathlib import Path
import re
import unittest
from typing import Any, Dict, List, Optional, Set, Tuple

import numpy as np
import pandas as pd
from fastapi import HTTPException, UploadFile
from fastapi.responses import JSONResponse
from openpyxl import Workbook

SOURCE = Path(os.environ.get('APP_SOURCE', 'main.py')).resolve()
# Load analytics definitions only, avoiding authentication, secrets and Mongo startup.
tree = ast.parse(SOURCE.read_text())
pure_end = next(n.lineno for n in tree.body if isinstance(n, ast.FunctionDef) and n.name == 'require_mongo_db')
nodes = []
for node in tree.body:
    if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef)) and (node.lineno < pure_end or node.name == 'upload_rolling_category'):
        node.decorator_list = []
        if node.name == 'upload_rolling_category':
            node.args.defaults = []
        nodes.append(node)
    elif isinstance(node, ast.Assign) and any(isinstance(t, ast.Name) and (t.id.startswith('_') or t.id == 'ALLOWED_EXTENSIONS') for t in node.targets):
        nodes.append(node)
ns = dict(globals(), logger=logging.getLogger('calculation-tests'))
exec(compile(ast.Module(body=nodes, type_ignores=[]), str(SOURCE), 'exec'), ns)


def analyze(rows, benchmark=None):
    frame = pd.DataFrame(rows)
    return ns['analyze_fund_dataset'](frame, ns['map_fund_columns'](frame), benchmark)


def workbook_result(values, formats, prefix=False, sheet_name=None):
    book = Workbook()
    sheet = book.active
    sheet.title = 'Returns'
    if prefix:
        sheet.append(['Research export'])
        sheet.append([])
    sheet.append(['Fund Name', 'Category', '1Y Return', '3Y Return', 'Sharpe'])
    for i, (value, fmt) in enumerate(zip(values, formats)):
        sheet.append([f'Fund {i}', 'Equity', value, 20, .6])
        sheet.cell(sheet.max_row, 3).number_format = fmt
        # Percentage normalization must not change dimensionless ratios.
        sheet.cell(sheet.max_row, 5).number_format = '0.00%'
    if sheet_name:
        book.create_sheet('Empty')
    stream = io.BytesIO()
    book.save(stream)
    parsed = ns['smart_read_sheet'](stream.getvalue(), 'returns.xlsx', sheet_name)
    return ns['analyze_fund_dataset'](parsed['df'], ns['map_fund_columns'](parsed['df'])), parsed


class CalculationTests(unittest.TestCase):
    def setUp(self):
        self.rows = [
            {'Fund Name': 'Atlas', 'Category': 'Equity', 'Sharpe': 1.5, 'AUM': 100, 'Information Ratio': .6, 'Treynor': 15, '1Y Return': 12, '3Y Return': 18},
            {'Fund Name': 'Birch', 'Category': 'Equity', 'Sharpe': .5, 'AUM': 200, 'Information Ratio': .2, 'Treynor': 5, '1Y Return': 6},
            {'Fund Name': 'Cedar', 'Category': 'Debt', 'Sharpe': -.5, 'AUM': 50, 'Information Ratio': -.2, 'Treynor': -5},
            {'Fund Name': 'Dune', 'Category': 'Debt', 'Sharpe': 'NA', 'AUM': 25},
        ]

    def test_summary_ranks_categories_and_distribution(self):
        data = analyze(self.rows)
        expected = dict(total_funds=4, valid_sharpe_count=3, avg_sharpe=.5,
                        median_sharpe=.5, std_sharpe=.82, total_aum=375,
                        above_avg_count=2, below_avg_count=1, na_count=1,
                        outperformance_rate=66.7, avg_info_ratio=.2, avg_treynor=5)
        for key, value in expected.items():
            with self.subTest(metric=key):
                self.assertAlmostEqual(data['summary'][key], value)
        self.assertEqual([f['rank'] for f in data['funds']], [1, 2, 3, None])
        self.assertEqual(data['category_stats']['Equity']['avg_sharpe'], 1)
        self.assertEqual(sum(b['count'] for b in data['distribution']), 3)
        self.assertEqual(data['funds'][0]['rolling_avg'], 15)

    def test_custom_benchmark(self):
        data = analyze(self.rows, 1)
        self.assertEqual(data['summary']['above_avg_count'], 1)
        self.assertEqual(data['funds'][0]['sharpe_diff'], .5)

    def test_empty_missing_and_nonfinite_values(self):
        self.assertEqual(ns['compute_records_analysis']([])['summary']['total_funds'], 0)
        self.assertEqual(analyze([self.rows[-1]])['summary']['valid_sharpe_count'], 0)
        for value in ['NA', '--', float('nan'), float('inf'), None]:
            self.assertIsNone(ns['clean_val'](value))
        self.assertEqual(ns['clean_val']('1,234.50%'), 1234.5)

    def test_information_ratio_does_not_match_first_nav_date(self):
        for header in ['Information Ratio', 'Info Ratio', 'IR', 'IR (3Y)']:
            frame = pd.DataFrame(columns=['Fund Name', 'First NAV Date', header, 'Sharpe'])
            with self.subTest(header=header):
                self.assertEqual(ns['map_fund_columns'](frame)['info_ratio'], header)
        self.assertIsNone(ns['map_fund_columns'](pd.DataFrame(columns=['Fund Name', 'First NAV Date']))['info_ratio'])

    def test_excel_percentages_preserve_units_and_ratios(self):
        data, _ = workbook_result([.12, -.03, 0, 12, .12, '12%'],
                                  ['0.00%', '0.0%', '0%', 'General', 'General', 'General'])
        self.assertEqual([f['rolling_1y'] for f in data['funds']], [12, -3, 0, 12, .12, 12])
        self.assertTrue(all(f['sharpe'] == .6 for f in data['funds']))

    def test_percentages_with_preamble_and_selected_sheet(self):
        data, parsed = workbook_result([.12], ['0.00%'], prefix=True, sheet_name='Returns')
        self.assertEqual(parsed['header_row'], 2)
        self.assertEqual(data['funds'][0]['rolling_1y'], 12)

    def test_legacy_xls_percentage_formats(self):
        fixture = SOURCE.parent / 'tests/fixtures/percentage_returns.xls'
        parsed = ns['smart_read_sheet'](fixture.read_bytes(), 'returns.xls')
        data = ns['analyze_fund_dataset'](parsed['df'], ns['map_fund_columns'](parsed['df']))
        self.assertEqual([f['rolling_1y'] for f in data['funds']], [12, -3, 12])
        self.assertTrue(all(f['sharpe'] == .6 for f in data['funds']))

    def test_literal_percent_formats_are_not_scaled(self):
        data, _ = workbook_result([12, 12], ['0.00"%"', r'0.00\%'])
        self.assertEqual([f['rolling_1y'] for f in data['funds']], [12, 12])

    def test_csv_percentages_are_not_guessed(self):
        parsed = ns['smart_read_sheet'](b'Fund Name,Category,1Y Return\nAtlas,Equity,0.12\nBirch,Equity,12%\n', 'returns.csv')
        data = ns['analyze_fund_dataset'](parsed['df'], ns['map_fund_columns'](parsed['df']))
        self.assertEqual([f['rolling_1y'] for f in data['funds']], [.12, 12])

    def test_sequential_upload_mean_uses_retained_horizons(self):
        current = analyze([{'Fund Name': 'Atlas Equity Regular Growth', 'Category': 'Equity', 'Sharpe': 1, '1Y Return': 10}])
        async def merge(funds, content):
            response = await ns['upload_rolling_category'](
                UploadFile(filename='returns.csv', file=io.BytesIO(content)),
                json.dumps(funds), None, {'uid': 'synthetic'})
            return json.loads(response.body)['data']['funds']
        funds = asyncio.run(merge(current['funds'], b'Fund Name,Category,3Y Return\nAtlas Equity Regular Growth,Equity,20\n'))
        self.assertEqual(funds[0]['rolling_1y'], 10)
        self.assertEqual(funds[0]['rolling_avg'], 15)
        funds = asyncio.run(merge(funds, b'Fund Name,Category,1Y Return,5Y Return\nAtlas Equity Regular Growth,Equity,0,-5\n'))
        self.assertEqual(funds[0]['rolling_avg'], 5)
        self.assertEqual(funds[0]['sharpe'], 1)

    def test_distinct_scheme_identities_never_merge(self):
        for query, candidate in [
            ('HDFC Mid Cap Fund Regular Growth', 'HDFC Large Cap Fund Regular Growth'),
            ('HDFC Large and Mid Cap Fund', 'HDFC Large Cap Fund'),
            ('Atlas Nifty 50 Index Fund', 'Atlas Nifty 500 Index Fund'),
            ('Atlas Equity Fund', 'Atlas Equity Index Fund'),
            ('Atlas Equity Direct Growth', 'Atlas Equity Regular Growth'),
        ]:
            with self.subTest(query=query, candidate=candidate):
                matcher = ns['build_rolling_matcher']([{'name': candidate}])
                self.assertIsNone(ns['match_rolling_record'](matcher, query)['record'])

    def test_equivalent_names_and_consumed_records(self):
        for query, candidate in [
            ('HDFC Midcap Fund Reg Gr', 'HDFC Mid Cap Regular Growth'),
            ('ICICI Pru Equity Fund', 'ICICI Prudential Equity Regular Growth'),
            ('Atlas Equity', 'Atlas Equity Regular Growth'),
        ]:
            with self.subTest(query=query):
                record = {'name': candidate}
                matcher = ns['build_rolling_matcher']([record])
                self.assertIs(ns['match_rolling_record'](matcher, query)['record'], record)
                self.assertIsNone(ns['match_rolling_record'](matcher, query, {id(record)})['record'])

    def test_different_schemes_remain_separate_in_merge(self):
        risk = analyze([{'Fund Name': 'HDFC Mid Cap Fund Regular Growth', 'Category': 'Mid Cap', 'Sharpe': 1}])
        rolling = analyze([{'Fund Name': 'HDFC Large Cap Fund Regular Growth', 'Category': 'Large Cap', '1Y Return': 99}])
        data = ns['merge_risk_and_rolling'](risk, rolling)
        self.assertEqual(len(data['funds']), 2)
        self.assertIsNone(data['funds'][0]['rolling_1y'])
        self.assertEqual(data['summary']['rolling_only_count'], 1)


if __name__ == '__main__':
    unittest.main(verbosity=2)
